from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference, LineChart

wb = load_workbook("sample2.xlsx")
ws = wb.active

# bar_value = Reference(ws, min_col=2, max_col=3, min_row=1, max_row=10)
# bar_chart = BarChart()
# bar_chart.add_data(bar_value, titles_from_data=True)
# bar_chart.title = "성적표"
# bar_chart.style = 10
# bar_chart.x_axis.title = "번호"
# bar_chart.y_axis.title = "점수"

# titles = Reference(ws, min_col=1, min_row=2, max_row=10)
# bar_chart.set_categories(titles)

# ws.add_chart(bar_chart, "E1")

chart_value = Reference(ws, min_col=2, max_col=3, min_row=1, max_row=10)
line_chart = LineChart()
line_chart.add_data(chart_value, titles_from_data=True)
line_chart.title = "성적표"
line_chart.style = 10
line_chart.x_axis.title = "번호"
line_chart.y_axis.title = "점수"


titles = Reference(ws, min_col=1, min_row=2, max_row=10)
line_chart.set_categories(titles)

ws.add_chart(line_chart, "E1")
wb.save("sample2_chart.xlsx")
